"""导出：Word / PDF / Markdown / 纯文本，可选“双语对照/仅译文”和“保留原文版式（页眉、页脚、页码与分页）”。"""
from __future__ import annotations

import html
import io
from dataclasses import dataclass


@dataclass
class Group:
    """一个导出单元：保留版式时对应原文的一页，否则整篇为一个单元。"""
    header: str
    footer: str
    paras: list[tuple[str, str]]  # (原文, 译文)


def make_groups(paras: list[dict], pages: list[dict] | None, preserve: bool) -> list[Group]:
    if not (preserve and pages):
        return [Group("", "", [(p["src"], p.get("tgt", "")) for p in paras])]
    groups = [Group(pg.get("header", ""), pg.get("footer", ""), []) for pg in pages]
    for p in paras:
        i = p.get("page")
        i = i if isinstance(i, int) and 0 <= i < len(groups) else (len(groups) - 1)
        groups[i].paras.append((p["src"], p.get("tgt", "")))
    return groups


FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pdf": ("application/pdf", "pdf"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "txt": ("text/plain; charset=utf-8", "txt"),
}


def export(fmt: str, title: str, paras: list[dict], pages: list[dict] | None,
           bilingual: bool, preserve: bool) -> bytes:
    groups = make_groups(paras, pages, preserve)
    fn = {"docx": to_docx, "pdf": to_pdf, "md": to_markdown, "txt": to_text}[fmt]
    return fn(groups, title, bilingual, preserve and bool(pages))


_MISSING = "（未翻译）"


def to_markdown(groups: list[Group], title: str, bilingual: bool, paged: bool) -> bytes:
    out = [] if paged else [f"# {title}", ""]
    for gi, g in enumerate(groups):
        if paged:
            if gi:
                out += ["", "---", ""]
            if g.header:
                out += [f"<div align=\"center\"><sub>{html.escape(g.header)}</sub></div>", ""]
        for src, tgt in g.paras:
            if bilingual:
                out += ["> " + src.replace("\n", "\n> "), ""]
            out += [tgt or _MISSING, ""]
        if paged and g.footer:
            out += [f"<div align=\"center\"><sub>{html.escape(g.footer)}</sub></div>"]
    return "\n".join(out).strip().encode("utf-8") + b"\n"


def to_text(groups: list[Group], title: str, bilingual: bool, paged: bool) -> bytes:
    out = [] if paged else [title, ""]
    for gi, g in enumerate(groups):
        if paged:
            if gi:
                out += ["", "\f"]  # 换页符：多数编辑器与打印机会在此分页
            if g.header:
                out += [g.header, ""]
        for src, tgt in g.paras:
            if bilingual:
                out += [src]
            out += [tgt or _MISSING, ""]
        if paged and g.footer:
            out += [g.footer.center(60).rstrip()]
    return "\n".join(out).strip().replace("\n", "\r\n").encode("utf-8-sig") + b"\r\n"


def to_docx(groups: list[Group], title: str, bilingual: bool, paged: bool) -> bytes:
    import docx
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing = 1.4
    style.paragraph_format.space_after = Pt(6)
    grey = RGBColor(0x80, 0x80, 0x80)

    def set_hf(part, text):
        part.is_linked_to_previous = False
        p = part.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = grey

    if not paged:
        d.add_heading(title, level=1)
    for gi, g in enumerate(groups):
        if paged:
            # 每个原文页对应一个 Word 节（section）：独立页眉页脚 + 从新页开始
            sec = d.sections[0] if gi == 0 else d.add_section(WD_SECTION.NEW_PAGE)
            set_hf(sec.header, g.header)
            set_hf(sec.footer, g.footer)
        for src, tgt in g.paras:
            if bilingual:
                r = d.add_paragraph().add_run(src)
                r.font.color.rgb = grey
                r.font.size = Pt(10)
            d.add_paragraph(tgt or _MISSING)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


_PDF_CSS = """
body { font-family: serif; font-size: 11pt; line-height: 1.6; color: #222; }
h1 { font-size: 17pt; text-align: center; margin-bottom: 14pt; }
p { margin: 0 0 8pt 0; text-align: justify; }
p.src { color: #7a7a7a; font-size: 9.5pt; margin-bottom: 3pt; }
"""


def to_pdf(groups: list[Group], title: str, bilingual: bool, paged: bool) -> bytes:
    import pymupdf

    a4 = pymupdf.paper_rect("a4")
    body = a4 + (64, 70, -64, -70)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    hf_pages = []  # 每个 PDF 页对应的 (页眉, 页脚)，最后统一绘制

    for gi, g in enumerate(groups):
        parts = [] if paged or gi else [f"<h1>{html.escape(title)}</h1>"]
        for src, tgt in g.paras:
            if bilingual:
                parts.append(f'<p class="src">{html.escape(src)}</p>')
            parts.append(f"<p>{html.escape(tgt or _MISSING)}</p>")
        if not parts:
            parts.append("<p></p>")
        story = pymupdf.Story(html="".join(parts), user_css=_PDF_CSS)
        more = 1
        while more:  # 一个原文页的译文放不下时自动续页
            dev = writer.begin_page(a4)
            more, _ = story.place(body)
            story.draw(dev)
            writer.end_page()
            hf_pages.append((g.header, g.footer) if paged else ("", ""))
    writer.close()

    doc = pymupdf.open("pdf", buf.getvalue())
    font = pymupdf.Font("cjk")
    for page, (header, footer) in zip(doc, hf_pages):
        page.wrap_contents()  # 隔离 Story 留下的坐标变换，否则页眉页脚会倒置
        for text, y in ((header, 36), (footer, a4.height - 50)):
            if text:
                tw = pymupdf.TextWriter(page.rect, color=(0.5, 0.5, 0.5))
                fs = 8.5
                width = font.text_length(text, fontsize=fs)
                if width > a4.width - 120:  # 过长时缩小字号以免超出页面
                    fs *= (a4.width - 120) / width
                    width = a4.width - 120
                tw.append(((a4.width - width) / 2, y + 8.5), text, font=font, fontsize=fs)
                tw.write_text(page)
    doc.set_metadata({"title": title, "creator": "灵译"})
    doc.subset_fonts()  # 只嵌入用到的字形，大幅减小文件体积
    return doc.tobytes(garbage=3, deflate=True)
